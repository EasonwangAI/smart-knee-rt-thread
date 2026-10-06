from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


PROJECT = Path(r"D:\RT-ThreadStudio\workspace\test_pro3_v8_dbui_kneeout")
OUT = next(p for p in PROJECT.iterdir() if p.is_dir() and p.name.startswith("Word"))
SOURCE_DIR = Path(r"C:\Users\39947\Documents\xwechat_files\wxid_3q4iendnkd1322_627b\msg\file\2026-08")


def strip_md(line: str) -> str:
    text = line.strip()
    text = re.sub(r"^#{1,6}\s+", "", text)
    text = re.sub(r"^>\s?", "", text)
    text = re.sub(r"^[-*+]\s+\[[ xX]\]\s+", "", text)
    text = re.sub(r"^[-*+]\s+", "", text)
    text = re.sub(r"^\d+[.)]\s+", "", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.strip().strip("|").strip()


for docx_path in sorted(OUT.glob("*.docx")):
    with zipfile.ZipFile(docx_path) as archive:
        corrupt = archive.testzip()
        required = {"word/document.xml", "word/styles.xml", "word/numbering.xml"}
        missing_parts = sorted(required.difference(archive.namelist()))

    doc = Document(docx_path)
    paragraphs = list(doc.paragraphs)
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                table_text.extend(p.text for p in cell.paragraphs)
    body_text = "\n".join([p.text for p in paragraphs] + table_text)
    headings = sum(1 for p in paragraphs if p.style.name.startswith("Heading"))
    numbered = sum(
        1 for p in paragraphs
        if p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None
    )

    source = SOURCE_DIR / f"{docx_path.stem}.md"
    missing_lines = []
    for raw in source.read_text(encoding="utf-8-sig").splitlines():
        clean = strip_md(raw)
        if (not clean or clean.startswith("```") or
                re.fullmatch(r":?-{3,}:?", clean.replace("|", ""))):
            continue
        # Table rows are split into cells in DOCX, so validate each nonempty cell.
        pieces = [p.strip() for p in clean.split("|") if p.strip()]
        for piece in pieces:
            if len(piece) >= 3 and piece not in body_text:
                missing_lines.append(piece)

    print(docx_path.name)
    print(f"  zip_corrupt={corrupt!r} missing_parts={missing_parts}")
    print(f"  paragraphs={len(paragraphs)} tables={len(doc.tables)} headings={headings} numbered={numbered}")
    print(f"  body_chars={len(body_text)} missing_source_fragments={len(missing_lines)}")
    for fragment in missing_lines[:10]:
        print(f"    MISSING: {fragment}")
