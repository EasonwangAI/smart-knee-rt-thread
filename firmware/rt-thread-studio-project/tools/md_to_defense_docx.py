from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable, List, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# Resolved design preset: compact_reference_guide.
# Named overrides: Microsoft YaHei for East Asian glyphs; editorial_cover title
# block; compact 9.5 pt body text inside genuinely tabular tables.
PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"
FONT_CODE = "Consolas"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
MUTED = "66717D"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
BORDER = "CBD5E1"
WHITE = "FFFFFF"
BLACK = "111827"


def set_run_font(run, name=FONT_LATIN, east_asia=FONT_CJK, size=None,
                 color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_fonts(style, size, color=BLACK, bold=False, name=FONT_LATIN,
                east_asia=FONT_CJK):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)


def set_style_spacing(style, before, after, line=1.25):
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def set_repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    tag = trpr.find(qn("w:tblHeader"))
    if tag is None:
        tag = OxmlElement("w:tblHeader")
        trpr.append(tag)
    tag.set(qn("w:val"), "true")


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for side, value in (("top", top), ("start", start),
                        ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: Sequence[int]):
    assert sum(widths) == CONTENT_WIDTH_DXA
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = table._tbl.tblPr

    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tblw.set(qn("w:type"), "dxa")

    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tblind.set(qn("w:type"), "dxa")

    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[idx]
            cell.width = Inches(width / 1440)
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_box(paragraph, fill=LIGHT_GRAY, left_color=BLUE,
                      all_borders=False):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")

    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    edges = ("top", "left", "bottom", "right") if all_borders else ("left",)
    for edge in edges:
        node = pbdr.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            pbdr.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "12" if edge == "left" else "4")
        node.set(qn("w:space"), "6")
        node.set(qn("w:color"), left_color if edge == "left" else BORDER)


def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MUTED)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    rpr.append(color)
    rpr.append(size)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def add_numbering_definition(doc, kind: str):
    numbering = doc.part.numbering_part.element
    abs_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abs_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    numfmt = OxmlElement("w:numFmt")
    numfmt.set(qn("w:val"), "decimal" if kind == "number" else "bullet")
    lvl.append(numfmt)
    lvltext = OxmlElement("w:lvlText")
    lvltext.set(qn("w:val"), "%1." if kind == "number" else ("□" if kind == "check" else "•"))
    lvl.append(lvltext)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    ppr.append(ind)
    lvl.append(ppr)
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    marker_font = "Microsoft YaHei" if kind == "check" else FONT_LATIN
    rfonts.set(qn("w:ascii"), marker_font)
    rfonts.set(qn("w:hAnsi"), marker_font)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)
    rpr.append(rfonts)
    lvl.append(rpr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.append(numpr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(numid)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT_LATIN)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)
    rpr.append(rfonts)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text, default_size=11, default_color=BLACK,
               default_bold=False):
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=default_size, color=default_color, bold=default_bold)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, color=default_color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=FONT_CODE, east_asia=FONT_CJK,
                         size=max(9, default_size - 1), color=DARK_BLUE)
            rpr = run._element.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EEF2F6")
            rpr.append(shd)
        else:
            lm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            add_hyperlink(paragraph, lm.group(1), lm.group(2))
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=default_size, color=default_color, bold=default_bold)


def configure_document(doc: Document, short_title: str):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    sec.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    style_fonts(normal, 11)
    set_style_spacing(normal, 0, 6, 1.25)
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style_fonts(style, size, color=color, bold=True)
        set_style_spacing(style, before, after, 1.1)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    for style_name in ("Header", "Footer"):
        style = doc.styles[style_name]
        style_fonts(style, 9, color=MUTED)
        set_style_spacing(style, 0, 0, 1.0)

    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    run = hp.add_run(short_title)
    set_run_font(run, size=9, color=MUTED, bold=True)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    add_page_field(fp)


def add_cover(doc: Document, title: str, subtitle: str):
    for _ in range(5):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("智能护膝项目 · 答辩材料")
    set_run_font(run, size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    set_run_font(run, size=28, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    run = p.add_run(subtitle)
    set_run_font(run, size=13, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(90)
    run = p.add_run("多模态感知 · 边缘智能 · 数据库校准 · 云端AI")
    set_run_font(run, size=10.5, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("Word 排版版")
    set_run_font(run, size=11, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026年8月")
    set_run_font(run, size=10, color=MUTED)

    doc.add_page_break()


def is_table_separator(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def split_table_row(line: str) -> List[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def choose_widths(rows: Sequence[Sequence[str]]) -> List[int]:
    n = len(rows[0])
    if n == 1:
        return [CONTENT_WIDTH_DXA]
    lengths = []
    for c in range(n):
        lengths.append(max(4, max(len(row[c]) if c < len(row) else 0 for row in rows)))
    total = sum(lengths)
    raw = [int(CONTENT_WIDTH_DXA * x / total) for x in lengths]
    minimum = 1200 if n <= 4 else 900
    raw = [max(minimum, x) for x in raw]
    scale = CONTENT_WIDTH_DXA / sum(raw)
    raw = [int(x * scale) for x in raw]
    raw[-1] += CONTENT_WIDTH_DXA - sum(raw)
    return raw


def add_markdown_table(doc: Document, rows: Sequence[Sequence[str]]):
    col_count = len(rows[0])
    normalized = [list(r) + [""] * (col_count - len(r)) for r in rows]
    table = doc.add_table(rows=len(normalized), cols=col_count)
    widths = choose_widths(normalized)
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(normalized):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            add_inline(p, value, default_size=9.5,
                       default_color=NAVY if r_idx == 0 else BLACK,
                       default_bold=(r_idx == 0))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, lines: Sequence[str], language: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.03)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    set_paragraph_box(p, fill="F7F9FC", left_color="94A3B8", all_borders=True)
    if language:
        run = p.add_run(f"{language.upper()}\n")
        set_run_font(run, name=FONT_CODE, east_asia=FONT_CJK,
                     size=8.5, color=BLUE, bold=True)
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, name=FONT_CODE, east_asia=FONT_CJK,
                     size=8.7, color=BLACK)
        if idx != len(lines) - 1:
            run.add_break()


def add_blockquote(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    set_paragraph_box(p, fill=LIGHT_GRAY, left_color=BLUE)
    add_inline(p, text, default_size=10.5, default_color=DARK_BLUE)


def convert_markdown(source: Path, output: Path, cover_subtitle: str):
    text = source.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    title = source.stem
    for line in lines:
        if line.startswith("# "):
            title = line[2:].strip()
            break

    doc = Document()
    configure_document(doc, title)
    doc.core_properties.title = title
    doc.core_properties.subject = "智能护膝项目答辩参考材料"
    doc.core_properties.author = "智能护膝项目组"
    doc.core_properties.keywords = "智能护膝, 答辩, 多模态感知, 边缘AI"
    add_cover(doc, title, cover_subtitle)

    i = 0
    first_h1_seen = False
    active_list_kind = None
    active_num_id = None
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            active_list_kind = None
            active_num_id = None
            i += 1
            continue

        if stripped.startswith("```"):
            language = stripped[3:].strip()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, block, language)
            active_list_kind = None
            active_num_id = None
            continue

        if ("|" in line and i + 1 < len(lines) and
                is_table_separator(lines[i + 1])):
            rows = [split_table_row(line)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(split_table_row(lines[i]))
                i += 1
            add_markdown_table(doc, rows)
            active_list_kind = None
            active_num_id = None
            continue

        hm = re.match(r"^(#{1,4})\s+(.+)$", line)
        if hm:
            level = len(hm.group(1))
            heading = hm.group(2).strip()
            if level == 1:
                i += 1
                continue
            style = "Heading 1" if level == 2 else ("Heading 2" if level == 3 else "Heading 3")
            p = doc.add_paragraph(style=style)
            p.paragraph_format.keep_with_next = True
            if style == "Heading 1" and not first_h1_seen:
                p.paragraph_format.space_before = Pt(0)
                first_h1_seen = True
            add_inline(p, heading,
                       default_size=16 if style == "Heading 1" else (13 if style == "Heading 2" else 12),
                       default_color=BLUE if style != "Heading 3" else DARK_BLUE,
                       default_bold=True)
            active_list_kind = None
            active_num_id = None
            i += 1
            continue

        if stripped.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                content = lines[i].strip()[1:].strip()
                if content:
                    quote_lines.append(content)
                i += 1
            add_blockquote(doc, " ".join(quote_lines))
            active_list_kind = None
            active_num_id = None
            continue

        checklist = re.match(r"^\s*-\s+\[([ xX])\]\s+(.+)$", line)
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        numbered = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
        if checklist or bullet or numbered:
            if checklist:
                kind, content = "check", checklist.group(2)
            elif bullet:
                kind, content = "bullet", bullet.group(1)
            else:
                kind, content = "number", numbered.group(1)
            if active_list_kind != kind or active_num_id is None:
                active_list_kind = kind
                active_num_id = add_numbering_definition(doc, kind)
            p = doc.add_paragraph()
            apply_num(p, active_num_id)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            add_inline(p, content)
            i += 1
            continue

        if re.fullmatch(r"-{3,}", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            ppr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:color"), BORDER)
            pbdr.append(bottom)
            ppr.append(pbdr)
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith("#") or nxt.startswith(">") or
                    nxt.startswith("```") or re.match(r"^\s*[-*+]\s+", lines[i]) or
                    re.match(r"^\s*\d+[.)]\s+", lines[i]) or
                    ("|" in lines[i] and i + 1 < len(lines) and is_table_separator(lines[i + 1]))):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines))
        active_list_kind = None
        active_num_id = None

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    sources = [
        (
            Path(r"C:\Users\39947\Documents\xwechat_files\wxid_3q4iendnkd1322_627b\msg\file\2026-08\智能护膝项目评委问答速查清单.md"),
            "70个高频问题 · 必背题标记 · 现场回答口径",
        ),
        (
            Path(r"C:\Users\39947\Documents\xwechat_files\wxid_3q4iendnkd1322_627b\msg\file\2026-08\智能护膝项目答辩全清单.md"),
            "系统架构 · 核心算法 · 演示流程 · 风险与检查清单",
        ),
    ]
    out_dir = Path(r"D:\RT-ThreadStudio\workspace\test_pro3_v8_dbui_kneeout\Word答辩材料")
    for source, subtitle in sources:
        if not source.exists():
            raise FileNotFoundError(source)
        output = out_dir / f"{source.stem}.docx"
        convert_markdown(source, output, subtitle)
        print(output)


if __name__ == "__main__":
    main()
